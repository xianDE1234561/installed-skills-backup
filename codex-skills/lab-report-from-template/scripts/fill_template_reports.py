import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SECTION_KEYS = {
    "purpose": "一、",
    "devices": "二、",
    "plan": "三、",
    "data": "四、",
    "conclusion": "五、",
}


def font_path():
    for candidate in [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
    ]:
        if Path(candidate).exists():
            return candidate
    return None


FONT_PATH = font_path()


def pil_font(size):
    return ImageFont.truetype(FONT_PATH, size=size) if FONT_PATH else ImageFont.load_default()


def set_run_font(run, size=10.5, bold=False, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def clear_para(paragraph):
    paragraph._p.clear_content()


def write_para(paragraph, text, size=10.5, bold=False, align=None, first_line=True):
    clear_para(paragraph)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_picture(paragraph, image_path, width_cm=14.3):
    clear_para(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def resolve_path(base_dir, value):
    path = Path(value)
    if path.is_absolute():
        return path
    from_config = (base_dir / path).resolve()
    if from_config.exists():
        return from_config
    return path.resolve()


def make_collage(report, config_dir, out_dir):
    figures = report.get("figures", [])[:3]
    if not figures:
        return None

    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / f"{Path(report['filename']).stem}_综合图.jpg"
    width, height = 1800, 820
    margin, gap, title_h = 34, 24, 48
    panel_w = (width - 2 * margin - 2 * gap) // 3
    panel_h = height - 2 * margin - title_h
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)
    title_font = pil_font(26)

    for i, item in enumerate(figures):
        label = item.get("label", f"图{i + 1}")
        image_path = resolve_path(config_dir, item["path"])
        x = margin + i * (panel_w + gap)
        y = margin
        draw.text((x + 6, y), label, fill=(30, 55, 85), font=title_font)
        box = (x, y + title_h, x + panel_w, y + title_h + panel_h)
        draw.rectangle(box, outline=(130, 150, 170), width=2)
        image = Image.open(image_path).convert("RGB")
        image.thumbnail((panel_w - 18, panel_h - 18))
        px = x + (panel_w - image.width) // 2
        py = y + title_h + (panel_h - image.height) // 2
        canvas.paste(image, (px, py))

    canvas.save(dest, quality=92)
    return dest


def find_body_cell(doc):
    if len(doc.tables) < 2:
        raise ValueError("Template must contain the main body table as the second table.")
    return doc.tables[1].rows[0].cells[0]


def find_section_indices(paragraphs):
    indices = {}
    for i, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text.startswith("实验名称"):
            indices["name"] = i
        for key, prefix in SECTION_KEYS.items():
            if text.startswith(prefix):
                indices[key] = i
    missing = [key for key in ["name", *SECTION_KEYS.keys()] if key not in indices]
    if missing:
        raise ValueError(f"Template body is missing section markers: {missing}")
    return indices


def section_slots(indices, key, total_paras):
    order = ["purpose", "devices", "plan", "data", "conclusion"]
    start = indices[key] + 1
    pos = order.index(key)
    end = indices[order[pos + 1]] if pos + 1 < len(order) else total_paras
    return list(range(start, end))


def fill_lines(paragraphs, slots, lines, max_lines=None):
    use_lines = lines if max_lines is None else lines[:max_lines]
    for slot, line in zip(slots, use_lines):
        write_para(paragraphs[slot], line)


def fill_report(template, report, config_dir, out_dir, collage_dir):
    doc = Document(str(template))
    body = find_body_cell(doc)
    paragraphs = body.paragraphs
    indices = find_section_indices(paragraphs)

    write_para(paragraphs[indices["name"]], f"实验名称    {report['experiment_name']}", size=12, first_line=False)

    sections = report.get("sections", {})
    for key in ["purpose", "devices", "plan"]:
        fill_lines(paragraphs, section_slots(indices, key, len(paragraphs)), sections.get(key, []))

    data_slots = section_slots(indices, "data", len(paragraphs))
    data_lines = sections.get("data", [])
    collage = make_collage(report, config_dir, collage_dir)
    if collage and len(data_slots) >= 4:
        before_count = min(9, max(1, len(data_slots) // 2 - 1), len(data_lines))
        fill_lines(paragraphs, data_slots[:before_count], data_lines[:before_count])
        add_picture(paragraphs[data_slots[before_count]], collage)
        if before_count + 1 < len(data_slots):
            write_para(
                paragraphs[data_slots[before_count + 1]],
                "图：原始记录、数据处理曲线及实验实时趋势",
                size=9,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                first_line=False,
            )
        remaining_slots = data_slots[before_count + 2 :]
        fill_lines(paragraphs, remaining_slots, data_lines[before_count:])
    else:
        fill_lines(paragraphs, data_slots, data_lines)

    fill_lines(paragraphs, section_slots(indices, "conclusion", len(paragraphs)), sections.get("conclusion", []))

    out_path = out_dir / report["filename"]
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def validate_output(path, template):
    doc = Document(str(path))
    template_doc = Document(str(template))
    sec = doc.sections[0]
    template_sec = template_doc.sections[0]
    same_page = (
        sec.page_width == template_sec.page_width
        and sec.page_height == template_sec.page_height
        and sec.top_margin == template_sec.top_margin
        and sec.left_margin == template_sec.left_margin
    )
    images = sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype)
    return same_page, images


def main():
    parser = argparse.ArgumentParser(description="Fill lab report DOCX templates from a JSON config.")
    parser.add_argument("--config", required=True, help="JSON config containing report content.")
    parser.add_argument("--template", required=True, help="DOCX template path.")
    parser.add_argument("--out", required=True, help="Output directory.")
    args = parser.parse_args()

    config_path = Path(args.config).resolve()
    config_dir = config_path.parent
    template = Path(args.template).resolve()
    out_dir = Path(args.out).resolve()
    collage_dir = out_dir / "_generated_figures"

    config = json.loads(config_path.read_text(encoding="utf-8-sig"))
    for report in config.get("reports", []):
        out_path = fill_report(template, report, config_dir, out_dir, collage_dir)
        same_page, images = validate_output(out_path, template)
        print(f"{out_path} | page_settings_match={same_page} | images={images}")


if __name__ == "__main__":
    main()
